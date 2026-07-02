# -*- coding: utf-8 -*-
"""文件解析 — Word(.docx) / PDF / PPT(.pptx) 提取文本

Word/PPT 直接提文字（免费），PDF 逐页转图片（走已有 OCR/VL 通道）。
"""
from __future__ import annotations

import base64
import io
import os
from typing import List, Tuple


def extract_docx(file_bytes: bytes) -> str:
    """从 .docx 提取所有段落文本。"""
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            paragraphs.append(text)
    # 也提表格内容
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                paragraphs.append(row_text)
    return "\n".join(paragraphs)


def extract_pptx(file_bytes: bytes) -> str:
    """从 .pptx 提取所有文本框内容。"""
    from pptx import Presentation
    prs = Presentation(io.BytesIO(file_bytes))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        lines.append(text)
        if lines:
            slides.append(f"--- 第{i}页 ---\n" + "\n".join(lines))
    return "\n\n".join(slides)


def extract_pdf(file_bytes: bytes) -> Tuple[str, List[str]]:
    """从 PDF 提取文本 + 逐页转为 base64 图片。

    返回: (text_content, [page_image_data_url, ...])
    """
    import fitz  # PyMuPDF

    doc = fitz.open(stream=file_bytes, filetype="pdf")
    text_lines = []
    images = []

    for page in doc:
        # 尝试提取原生文字
        page_text = page.get_text().strip()
        if page_text:
            text_lines.append(page_text)

        # 逐页渲染为图片（无论有没有文字都做，用于 VL 理解）
        pix = page.get_pixmap(dpi=150)
        img_bytes = pix.tobytes("png")
        data_url = "data:image/png;base64," + base64.b64encode(img_bytes).decode("ascii")
        images.append(data_url)

    doc.close()
    return "\n".join(text_lines), images


# ── 统一入口 ──────────────────────────────────────

def parse_document(file_bytes: bytes, filename: str) -> Tuple[str, List[str]]:
    """根据文件扩展名自动选择解析器。

    返回: (text_content, images_data_urls)
      - Word/PPT: text 有内容, images 为空
      - PDF:      text 有内容（如有原生文字）, images 有逐页截图
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext in (".docx", ".doc"):
        return extract_docx(file_bytes), []
    elif ext in (".pptx", ".ppt"):
        return extract_pptx(file_bytes), []
    elif ext == ".pdf":
        return extract_pdf(file_bytes)
    else:
        raise ValueError(f"不支持的文件格式: {ext}")
